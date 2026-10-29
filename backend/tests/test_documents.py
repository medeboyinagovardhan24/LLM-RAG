"""Tests for the Documents endpoints (upload, list, get, delete)."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path


def _docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("Word paragraph content for retrieval.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Status"
    table.cell(1, 1).text = "Approved"
    output = BytesIO()
    document.save(output)
    return output.getvalue()

def _upload(client, data: bytes, name: str = "doc.pdf", content_type: str = "application/pdf"):
    return client.post("/api/upload", files={"file": (name, data, content_type)})


# --- Upload ---------------------------------------------------------------

def test_upload_valid_pdf(client, sample_pdf):
    resp = _upload(client, sample_pdf)
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is True
    assert body["duplicate"] is False
    assert body["filename"] == "doc.pdf"
    assert body["pages"] == 1
    assert body["chunks"] >= 1


def test_upload_rejects_non_pdf_extension(client):
    resp = _upload(client, b"hello", name="notes.txt", content_type="text/plain")
    assert resp.status_code == 400
    assert "PDF" in resp.json()["detail"]


def test_upload_valid_docx_extracts_paragraphs_and_tables(client):
    resp = _upload(
        client,
        _docx_bytes(),
        name="notes.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is True
    assert body["filename"] == "notes.docx"
    assert body["pages"] is None
    assert body["chunks"] >= 2
    assert body["tables"] == 1


def test_document_record_does_not_invent_docx_page_count(client):
    response = _upload(
        client,
        _docx_bytes(),
        name="pagination-unknown.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert response.status_code == 200
    document = client.get("/api/documents").json()["documents"][0]
    assert document["filename"] == "pagination-unknown.docx"
    assert document["pages"] is None


def test_docx_extraction_preserves_multiple_sections_and_table_code(tmp_path):
    from app.config import AppConfig
    from app.services import read_docx_into_chunks
    from docx import Document

    document = Document()
    document.add_paragraph("Experiment No. 1: Text Processing")
    document.add_paragraph("First experiment explanation.")
    first_table = document.add_table(rows=2, cols=1)
    first_table.cell(0, 0).text = "import nltk"
    first_table.cell(1, 0).text = "tokenize text"
    document.add_paragraph("Experiment No. 2: Deep Neural Network for a Classification Task")
    document.add_paragraph("Program")
    second_table = document.add_table(rows=2, cols=1)
    second_table.cell(0, 0).text = "import tensorflow as tf"
    second_table.cell(1, 0).text = "model.fit(x_train, y_train)"

    source = tmp_path / "multi-section.docx"
    document.save(source)
    _, characters, chunks, metadatas, _ = read_docx_into_chunks(
        Path(source),
        source.name,
        "test-hash",
        AppConfig(),
    )

    assert characters > 0
    assert len(chunks) >= 3
    code_chunks = [
        chunk for chunk, metadata in zip(chunks, metadatas)
        if metadata["content_type"] == "table" and "tensorflow" in chunk
    ]
    assert len(code_chunks) == 1
    assert "Experiment No. 2: Deep Neural Network for a Classification Task" in code_chunks[0]
    assert "model.fit(x_train, y_train)" in code_chunks[0]


def test_uploaded_docx_code_is_indexed_for_search(client, collection, tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Experiment No. 2: Deep Neural Network for a Classification Task")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "import tensorflow as tf; model.fit(x_train, y_train)"
    source = tmp_path / "code-manual.docx"
    document.save(source)

    response = _upload(
        client,
        source.read_bytes(),
        name=source.name,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert response.status_code == 200
    indexed = collection.get(
        where={"filename": source.name},
        include=["documents", "metadatas"],
    )
    assert any("model.fit(x_train, y_train)" in text for text in indexed["documents"])
    assert any(metadata["content_type"] == "table" for metadata in indexed["metadatas"])


def test_upload_rejects_invalid_docx_content(client):
    resp = _upload(
        client,
        b"not a Word document",
        name="notes.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert resp.status_code == 400
    assert "valid DOCX" in resp.json()["detail"]


def test_upload_rejects_empty_file(client):
    resp = _upload(client, b"")
    assert resp.status_code == 400
    assert "empty" in resp.json()["detail"].lower()


def test_upload_rejects_fake_pdf_content(client):
    resp = _upload(client, b"this is not really a pdf")
    assert resp.status_code == 400
    assert "valid PDF" in resp.json()["detail"]


def test_upload_rejects_unreadable_pdf(client):
    # Correct magic bytes but not a parseable document.
    resp = _upload(client, b"%PDF-1.4\nnot a real structure")
    assert resp.status_code == 422


def test_upload_rejects_oversized(client, monkeypatch):
    from app.config import get_settings

    monkeypatch.setenv("MAX_UPLOAD_MB", "1")
    get_settings.cache_clear()

    big = b"%PDF-1.4\n" + b"0" * (2 * 1024 * 1024)  # ~2 MB, valid magic bytes
    resp = _upload(client, big, name="big.pdf")
    assert resp.status_code == 413
    assert "limit" in resp.json()["detail"].lower()


def test_upload_detects_duplicate(client, sample_pdf):
    first = _upload(client, sample_pdf)
    assert first.json()["success"] is True

    second = _upload(client, sample_pdf)
    assert second.status_code == 200
    assert second.json()["duplicate"] is True


def test_upload_normalizes_windows_path_in_filename(client, sample_pdf):
    response = _upload(
        client,
        sample_pdf,
        name=r"..\..\safe.pdf",
    )

    assert response.status_code == 200
    assert response.json()["filename"] == "safe.pdf"


def test_upload_does_not_overwrite_existing_file_with_same_name(client, sample_pdf, data_dir):
    first = _upload(client, sample_pdf, name="same.pdf")
    assert first.status_code == 200
    original = (data_dir / "uploads" / "same.pdf").read_bytes()

    response = _upload(client, sample_pdf + b"\n", name="same.pdf")

    assert response.status_code == 409
    assert (data_dir / "uploads" / "same.pdf").read_bytes() == original
    assert client.get("/api/documents").json()["count"] == 1


# --- List / get / delete --------------------------------------------------

def test_list_documents(client, sample_pdf):
    assert client.get("/api/documents").json()["count"] == 0

    _upload(client, sample_pdf, name="rag.pdf")

    body = client.get("/api/documents").json()
    assert body["count"] == 1
    assert body["documents"][0]["filename"] == "rag.pdf"
    assert body["documents"][0]["chunks"] >= 1


def test_get_single_document(client, sample_pdf):
    _upload(client, sample_pdf, name="rag.pdf")
    resp = client.get("/api/documents/rag.pdf")
    assert resp.status_code == 200
    assert resp.json()["document"]["filename"] == "rag.pdf"


def test_get_missing_document_404(client):
    resp = client.get("/api/documents/nope.pdf")
    assert resp.status_code == 404


def test_delete_document(client, sample_pdf):
    _upload(client, sample_pdf, name="rag.pdf")

    resp = client.delete("/api/documents/rag.pdf")
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is True
    assert body["chunks_deleted"] >= 1

    assert client.get("/api/documents").json()["count"] == 0


def test_delete_missing_document_404(client):
    resp = client.delete("/api/documents/nope.pdf")
    assert resp.status_code == 404
